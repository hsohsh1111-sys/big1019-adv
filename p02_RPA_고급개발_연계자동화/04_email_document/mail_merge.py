"""메일 머지: 엑셀 명단 → 개인별 수료증(Word) 생성 → 첨부 메일 발송 → 발송 결과 엑셀.

실행:
  uvicorn web:app --reload   # 브라우저 → http://127.0.0.1:8000 (테스트 메일 서버도 자동 실행)
  python mail_server.py      (터미널 1)  →  python mail_merge.py   (터미널 2)
  또는 한 번에:  python mail_merge.py --demo   (테스트 서버를 내부에서 같이 실행)
"""
import html
import smtplib
import sys
import time
from datetime import date
from email.message import EmailMessage
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook, load_workbook

BASE = Path(__file__).parent
MEMBERS = BASE / "members.xlsx"
OUT = BASE / "output"
SMTP_HOST, SMTP_PORT = "127.0.0.1", 1025
SENDER = "edu-center@example.com"


def ensure_members():
    if MEMBERS.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.append(["이름", "이메일", "과정", "점수"])
    for row in [("김민수", "minsu@example.com", "생성형 AI 활용 개발", 92),
                ("이서연", "seoyeon@example.com", "RPA 고급 개발", 88),
                ("박지훈", "잘못된주소", "MLOps", 95),
                ("최유진", "yujin@example.com", "보안 및 DevOps", 79)]:
        ws.append(row)
    wb.save(MEMBERS)
    print(f"샘플 명단 생성: {MEMBERS.name}")


def read_members():
    ensure_members()
    return [list(r) for r in load_workbook(MEMBERS).active.iter_rows(min_row=2, values_only=True)]


def make_certificate(name, course, score):
    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("수 료 증")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(32), RGBColor(0x1F, 0x38, 0x64)
    doc.add_paragraph()
    table = doc.add_table(rows=3, cols=2, style="Table Grid")
    for i, (k, v) in enumerate([("성명", name), ("과정", course), ("점수", f"{score}점")]):
        table.cell(i, 0).text, table.cell(i, 1).text = k, v
    doc.add_paragraph()
    body = doc.add_paragraph(f"위 사람은 「{course}」 과정을 성실히 이수하였기에 이 증서를 수여합니다.")
    body.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer = doc.add_paragraph(f"\n{date.today():%Y년 %m월 %d일}\n\nCodeLab 교육센터장")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    path = OUT / f"수료증_{name}.docx"
    doc.save(path)
    return path


def make_mail(name, email, course, attachment):
    msg = EmailMessage()
    msg["From"], msg["To"], msg["Subject"] = SENDER, email, f"[CodeLab] {name}님 수료증 발급 안내"
    msg.set_content(f"{name}님, 「{course}」 과정 수료를 축하드립니다.\n첨부된 수료증을 확인해 주세요.")
    msg.add_alternative(f"<p><b>{html.escape(name)}</b>님, 「{html.escape(course)}」 과정 수료를 축하드립니다.</p>"
                        "<p>첨부된 수료증을 확인해 주세요.</p>", subtype="html")
    msg.add_attachment(attachment.read_bytes(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.wordprocessingml.document", filename=attachment.name)
    return msg


def send_all():
    """명단 전체에 수료증 메일을 보내고 건별 결과 목록을 반환. 메일 서버가 없으면 ConnectionError."""
    rows = read_members()
    OUT.mkdir(exist_ok=True)
    log = Workbook()
    ls = log.active
    ls.append(["이름", "이메일", "결과", "비고"])
    results = []

    try:
        smtp = smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=5)
    except OSError as e:
        raise ConnectionError("메일 서버에 연결할 수 없습니다. 먼저 `python mail_server.py` 를 실행하세요.") from e

    with smtp:
        for name, email, course, score in rows:
            try:
                if "@" not in str(email):
                    raise ValueError("이메일 형식 오류")
                cert = make_certificate(name, course, score)
                smtp.send_message(make_mail(name, email, course, cert))
                row = [name, email, "발송", cert.name]
            except (ValueError, smtplib.SMTPException) as e:  # 한 건 실패해도 계속
                row = [name, email, "실패", str(e)]
            ls.append(row)
            results.append(row)
            print(("✔" if row[2] == "발송" else "✘"), name, email, row[3])
            time.sleep(0.3)  # 대량 발송 시 스팸 차단 방지용 간격

    log.save(OUT / "send_result.xlsx")
    return results


def main():
    server = None
    if "--demo" in sys.argv:
        import mail_server
        server = mail_server.start()
        print("내장 테스트 메일 서버 시작")
    try:
        send_all()
        print(f"\n결과: {OUT / 'send_result.xlsx'} / 수료증: {OUT}")
    except ConnectionError as e:
        sys.exit(str(e) + " (또는 --demo 옵션 사용)")
    finally:
        if server:
            time.sleep(0.5)
            server.stop()


if __name__ == "__main__":
    main()
