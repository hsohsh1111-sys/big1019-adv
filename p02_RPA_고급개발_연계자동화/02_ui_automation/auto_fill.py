"""엑셀 명단을 읽어 웹 신청 폼에 자동 입력하는 UI 자동화 (Playwright).

실행:
  uvicorn web:app --reload     # 브라우저 대시보드 → http://127.0.0.1:8000 에서 [봇 실행]
  python auto_fill.py          # 화면 없이(headless) 실행
  python auto_fill.py --show   # 브라우저가 움직이는 모습 보기
결과: output/result.xlsx (건별 성공/실패), output/screenshot.png

브라우저: PC에 설치된 Edge/Chrome을 사용하고, 없으면 `playwright install chromium` 으로 설치한 브라우저 사용.
"""
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from playwright.sync_api import Error as PWError
from playwright.sync_api import sync_playwright

BASE = Path(__file__).parent
FORM = (BASE / "form.html").resolve().as_uri()
DATA = BASE / "applicants.xlsx"
OUT = BASE / "output"


def ensure_sample_data():
    if DATA.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.append(["이름", "이메일", "과정"])
    for row in [("김민수", "minsu@example.com", "생성형 AI"), ("이서연", "seoyeon@example.com", "RPA"),
                ("박지훈", "jihoon@example.com", "MLOps"), ("최유진", "yujin@example.com", "없는과정"),
                ("정하늘", "haneul@example.com", "RPA")]:
        ws.append(row)
    wb.save(DATA)
    print(f"샘플 명단 생성: {DATA.name}")


def read_rows(path):
    return [list(r) for r in load_workbook(path).active.iter_rows(min_row=2, values_only=True)]


def launch(p, headless):
    for channel in ("msedge", "chrome", None):  # 설치된 브라우저부터 시도
        try:
            return p.chromium.launch(channel=channel, headless=headless, slow_mo=0 if headless else 300)
        except PWError:
            continue
    sys.exit("브라우저를 찾지 못했습니다. `playwright install chromium` 을 실행하세요.")


def fill_one(page, name, email, course):
    page.get_by_label("이름").fill(name)
    page.get_by_label("이메일").fill(email)
    page.get_by_label("과정").select_option(course)  # 목록에 없는 값이면 예외 발생
    page.get_by_label("개인정보 수집 동의").check()
    page.get_by_role("button", name="신청").click()
    msg = page.locator("#msg.ok")
    msg.wait_for(timeout=5000)  # 처리 완료 메시지가 나타날 때까지 기다림 (sleep 사용 X)
    return msg.inner_text()


def main():
    ensure_sample_data()
    rows = read_rows(DATA)
    OUT.mkdir(exist_ok=True)
    results = Workbook()
    rs = results.active
    rs.append(["이름", "이메일", "과정", "결과", "메시지"])

    with sync_playwright() as p:
        browser = launch(p, headless="--show" not in sys.argv)
        page = browser.new_page()
        page.set_default_timeout(5000)  # 요소를 못 찾으면 5초 후 실패 처리
        page.goto(FORM)
        for name, email, course in rows:
            try:
                message = fill_one(page, name, email, course)
                rs.append([name, email, course, "성공", message])
                print(f"✔ {message}")
            except PWError as e:  # 한 건 실패해도 다음 사람 계속 진행
                rs.append([name, email, course, "실패", str(e).splitlines()[0][:80]])
                print(f"✘ {name}: 입력 실패 → 결과 파일에 기록")
                page.locator("#apply").evaluate("form => form.reset()")  # 입력하던 값 지우고 다음 건으로
        page.screenshot(path=OUT / "screenshot.png", full_page=True)
        browser.close()

    results.save(OUT / "result.xlsx")
    ok = sum(1 for r in rs.iter_rows(min_row=2, values_only=True) if r[3] == "성공")
    print(f"\n완료: 성공 {ok} / 전체 {len(rows)} → {OUT / 'result.xlsx'}, screenshot.png")


if __name__ == "__main__":
    main()
